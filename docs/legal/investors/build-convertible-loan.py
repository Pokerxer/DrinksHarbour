from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
base=Path('/Users/mac/Documents/drinksharbour/docs/legal/investors')
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(6)
 s.paragraph_format.line_spacing=1.05
normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
normal.paragraph_format.widow_control=True
doc.styles['Title'].font.size=Pt(24)
doc.styles['Heading 1'].font.size=Pt(14)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 1'].paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0];footer.alignment=2
r=footer.add_run('DrinksHarbour  |  Draft  |  ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
lines=(base/'drinksharbour-convertible-loan.md').read_text().splitlines()
i=0
schedule_mode=False
while i<len(lines):
 line=lines[i]
 if line=='---PAGE---':doc.add_page_break()
 elif line.startswith('# '):doc.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):
  if line.startswith('## Schedule 2'):schedule_mode=False
  doc.add_paragraph(line[3:],'Heading 1')
 elif line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].startswith('|'):
   vals=[x.strip() for x in lines[i].strip('|').split('|')]
   if not all(set(v)<=set('- :') for v in vals):rows.append(vals)
   i+=1
  table=doc.add_table(rows=0,cols=len(rows[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER
  table.autofit=False
  widths=[2.05,1.05,1.05,1.05,1.35] if len(rows[0])==5 else [2.55,1.25,2.75]
  for col,w in zip(table.columns,widths):col.width=Inches(w)
  for ri,vals in enumerate(rows):
   cells=table.add_row().cells
   for cell,w,val in zip(cells,widths,vals):
    cell.width=Inches(w);cell.text=val;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr=cell._tc.get_or_add_tcPr()
    borders=OxmlElement('w:tcBorders')
    for side in ['top','left','bottom','right']:
     el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    tcpr.append(borders)
    margins=OxmlElement('w:tcMar')
    for side in ['top','left','bottom','right']:
     el=OxmlElement('w:'+side);el.set(qn('w:w'),'55' if side in ['top','bottom'] else '85');el.set(qn('w:type'),'dxa');margins.append(el)
    tcpr.append(margins)
    if ri==0:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8E8E8');tcpr.append(sh)
    for p in cell.paragraphs:
     p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
     for run in p.runs:run.font.size=Pt(9);run.bold=(ri==0)
   if ri==0:
    rep=OxmlElement('w:tblHeader');table.rows[-1]._tr.get_or_add_trPr().append(rep)
  spacer=doc.add_paragraph()
  spacer.paragraph_format.space_after=Pt(0)
  spacer.paragraph_format.space_before=Pt(0)
  spacer.paragraph_format.line_spacing=Pt(3)
  continue
 elif line.strip():
  p=doc.add_paragraph(line)
  if schedule_mode:
   p.paragraph_format.space_after=Pt(4)
   p.paragraph_format.line_spacing=1.0
   for r in p.runs:r.font.size=Pt(9.5)
  if line=='DRAFT FOR NEGOTIATION AND LEGAL REVIEW':
   for r in p.runs:r.bold=True;r.font.size=Pt(9)
 i+=1
for root in [doc.styles.element, doc.element]:
 for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='DrinksHarbour Convertible Loan Agreement'
doc.core_properties.subject='Convertible bridge loan review draft'
doc.core_properties.author='DrinksHarbour'
doc.save(base/'DrinksHarbour Convertible Loan Agreement.docx')
print('Created Word agreement; source words:', len((base/'drinksharbour-convertible-loan.md').read_text().split()))
