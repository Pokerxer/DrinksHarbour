"""Build the three current review documents from their Markdown sources."""
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

BASE = Path(__file__).resolve().parent
FILES = [
    ('drinksharbour-nigerian-loan.md', 'DrinksHarbour Nigerian Advertising Loan Agreement.docx'),
    ('nigerian-company-valuation.md', 'DrinksHarbour Nigerian Company Valuation.docx'),
    ('nigerian-funding-plan.md', 'DrinksHarbour Nigerian Funding Plan.docx'),
]


def inline(paragraph, value):
    for index, part in enumerate(re.split(r'\*\*(.*?)\*\*', value)):
        run = paragraph.add_run(part)
        run.bold = index % 2 == 1


def table(document, lines):
    rows = [[v.strip() for v in line.strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(set(v) <= set('- :') for v in r)]
    count = len(rows[0])
    widths = {2: [2.25, 4.42], 3: [2.8, 1.3, 2.57], 4: [2.75, 1.3, 1.3, 1.32]}[count]
    if count == 3 and rows[0][0] == 'Buyer rights':
        widths = [3.1, 1.3, 2.27]
    t = document.add_table(rows=0, cols=count)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for index, row in enumerate(rows):
        cells = t.add_row().cells
        for cell, width, value in zip(cells, widths, row):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            inline(cell.paragraphs[0], value)
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'bottom', 'right'):
                edge = OxmlElement('w:' + side)
                for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    edge.set(qn('w:' + key), val)
                borders.append(edge)
            props.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                edge = OxmlElement('w:' + side)
                edge.set(qn('w:w'), '75' if side in ('top', 'bottom') else '95')
                edge.set(qn('w:type'), 'dxa')
                margins.append(edge)
            props.append(margins)
            if index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8E8E8'); props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    if index == 0:
                        run.bold = True
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); t.rows[-1]._tr.get_or_add_trPr().append(repeat)
        no_split = OxmlElement('w:cantSplit'); t.rows[-1]._tr.get_or_add_trPr().append(no_split)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = Pt(3)


for source, target in FILES:
    if len(sys.argv) > 1 and source not in sys.argv[1:]:
        continue
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
    sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.8); sec.right_margin = Inches(.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Calibri'; style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.05
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.widow_control = True
    doc.styles['Title'].font.size = Pt(23)
    doc.styles['Heading 1'].font.size = Pt(14)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
    doc.styles['Heading 1'].paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]; footer.alignment = 2
    footer.add_run('DrinksHarbour  |  Review draft  |  ').font.size = Pt(8)
    page = OxmlElement('w:fldSimple'); page.set(qn('w:instr'), 'PAGE'); footer._p.append(page)
    lines = (BASE / source).read_text().splitlines()
    i = 0
    page_break_pending = False
    while i < len(lines):
        line = lines[i]
        if line == '---PAGE---':
            page_break_pending = True
        elif line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            paragraph = doc.add_paragraph(line[3:], 'Heading 1')
            if page_break_pending:
                paragraph.paragraph_format.page_break_before = True
                page_break_pending = False
        elif line.startswith('|'):
            data = []
            while i < len(lines) and lines[i].startswith('|'):
                data.append(lines[i]); i += 1
            table(doc, data)
            continue
        elif line.strip():
            inline(doc.add_paragraph(), line)
        i += 1
    for root in (doc.styles.element, doc.element):
        for border in list(root.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.core_properties.title = lines[0][2:]
    doc.core_properties.author = 'DrinksHarbour'
    doc.core_properties.subject = 'Proposed terms for review, not executed'
    doc.save(BASE / target)
    print('Created:', target)
