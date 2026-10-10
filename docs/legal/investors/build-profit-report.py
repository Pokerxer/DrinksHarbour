from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

base = Path(__file__).resolve().parent
source = base / 'drinksharbour-revenue-and-profit.md'
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.8)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    style = document.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.03
document.styles['Normal'].font.size = Pt(10.5)
document.styles['Normal'].paragraph_format.widow_control = True
document.styles['Title'].font.size = Pt(23)
document.styles['Heading 1'].font.size = Pt(14)
document.styles['Heading 1'].paragraph_format.space_before = Pt(10)
document.styles['Heading 1'].paragraph_format.keep_with_next = True
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('DrinksHarbour  |  Revenue and profit  |  ').font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

lines = source.read_text().splitlines()
index, pending_break = 0, False
while index < len(lines):
    line = lines[index]
    if line == '---PAGE---':
        pending_break = True
    elif line.startswith('# '):
        document.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        paragraph = document.add_paragraph(line[3:], 'Heading 1')
        if pending_break:
            paragraph.paragraph_format.page_break_before = True
            pending_break = False
    elif line.startswith('|'):
        rows = []
        while index < len(lines) and lines[index].startswith('|'):
            values = [value.strip() for value in lines[index].strip('|').split('|')]
            if not all(set(value) <= set('- :') for value in values):
                rows.append(values)
            index += 1
        count = len(rows[0])
        widths = {2: [4.2, 2.45], 3: [2.85, 1.8, 2],
                  4: [2.55, 1.35, 1.35, 1.4], 5: [2.25, 1.1, 1.1, 1.1, 1.1]}[count]
        table = document.add_table(rows=0, cols=count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for row_index, values in enumerate(rows):
            row = table.add_row()
            for column_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width, cell.text = Inches(width), value
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for side in ('top', 'left', 'bottom', 'right'):
                    edge = OxmlElement('w:' + side)
                    for key, setting in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
                        edge.set(qn('w:' + key), setting)
                    borders.append(edge)
                properties.append(borders)
                padding = OxmlElement('w:tcMar')
                for side in ('top', 'left', 'bottom', 'right'):
                    margin = OxmlElement('w:' + side)
                    margin.set(qn('w:w'), '70' if side in ('top', 'bottom') else '95')
                    margin.set(qn('w:type'), 'dxa')
                    padding.append(margin)
                properties.append(padding)
                if row_index == 0:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'E8E8E8')
                    properties.append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    if column_index and count >= 4:
                        paragraph.alignment = 2
                    for run in paragraph.runs:
                        run.font.size, run.bold = Pt(9), row_index == 0
            if row_index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        document.add_paragraph().paragraph_format.space_after = Pt(2)
        continue
    elif line.strip():
        document.add_paragraph(line)
    index += 1

for root in (document.styles.element, document.element):
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
document.core_properties.title = 'How DrinksHarbour Earns Revenue and Builds Profit'
document.core_properties.subject = 'Business model operating economics and investor returns'
document.core_properties.author = 'DrinksHarbour'
output = base / 'DrinksHarbour Revenue and Profit Explained.docx'
document.save(output)
print(f'Created {output.name}; source words: {len(source.read_text().split())}')
