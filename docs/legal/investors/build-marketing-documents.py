"""Build the separate zero-equity agreement and concise funding plan."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

BASE = Path(__file__).resolve().parent


def build(source, output, title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.8)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.05
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.widow_control = True
    doc.styles['Title'].font.size = Pt(24)
    heading = doc.styles['Heading 1']
    heading.font.size = Pt(14)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('DrinksHarbour  |  Draft  |  ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    lines = (BASE / source).read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line == '---PAGE---':
            doc.add_page_break()
        elif line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            doc.add_paragraph(line[3:], 'Heading 1')
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                values = [v.strip() for v in lines[i].strip('|').split('|')]
                if not all(set(v) <= set('- :') for v in values):
                    rows.append(values)
                i += 1
            widths = [2.1, 4.57] if len(rows[0]) == 2 else [2.8, 1.29, 1.29, 1.29]
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            for col, width in zip(table.columns, widths):
                col.width = Inches(width)
            for ri, values in enumerate(rows):
                row = table.add_row()
                for ci, (cell, width, value) in enumerate(zip(row.cells, widths, values)):
                    cell.width = Inches(width)
                    cell.text = value
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    props = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement('w:tcBorders')
                    for side in ('top', 'left', 'bottom', 'right'):
                        el = OxmlElement('w:' + side)
                        for key, val in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
                            el.set(qn('w:' + key), val)
                        borders.append(el)
                    props.append(borders)
                    margins = OxmlElement('w:tcMar')
                    for side in ('top', 'left', 'bottom', 'right'):
                        el = OxmlElement('w:' + side)
                        el.set(qn('w:w'), '65' if side in ('top', 'bottom') else '90')
                        el.set(qn('w:type'), 'dxa')
                        margins.append(el)
                    props.append(margins)
                    if ri == 0:
                        shade = OxmlElement('w:shd')
                        shade.set(qn('w:fill'), 'E8E8E8')
                        props.append(shade)
                    for p in cell.paragraphs:
                        p.alignment = 0 if ci == 0 or len(values) == 2 else 2
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.space_before = Pt(2)
                        for run in p.runs:
                            run.font.size = Pt(9)
                            run.bold = ri == 0
                if ri == 0:
                    repeat = OxmlElement('w:tblHeader')
                    row._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        elif line.strip():
            paragraph = doc.add_paragraph(line)
            if line == 'DRAFT FOR NEGOTIATION AND LEGAL REVIEW':
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(9)
        i += 1
    for root in (doc.styles.element, doc.element):
        for border in list(root.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.core_properties.title = title
    doc.core_properties.author = 'DrinksHarbour'
    doc.core_properties.subject = 'Non-convertible advertising funding review draft'
    doc.save(BASE / output)
    print(output)


build('drinksharbour-marketing-loan.md', 'DrinksHarbour Advertising Loan Agreement.docx',
      'DrinksHarbour Advertising Loan Agreement')
build('marketing-funding-summary.md', 'DrinksHarbour Advertising Funding Plan.docx',
      'DrinksHarbour Advertising Funding Plan')
